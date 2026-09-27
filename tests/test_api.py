"""Real PostgreSQL API tests. No model requests are sent by this suite."""
import json
import uuid
import asyncio
from copy import deepcopy
from io import BytesIO
from psycopg.types.json import Jsonb
from types import SimpleNamespace
import psycopg
import pytest
from fastapi.testclient import TestClient
from branch_agent.app import create_app, BASE
from branch_agent.storage import Store
from branch_agent.records import new_record


class IdleEngine:
    def __init__(self):self.imports=[]
    def import_source(self,p,c,text,name,*,start_step1=False,step1_block_reason=None):
        self.imports.append({'project_id':p,'conversation_id':c,'text':text,'name':name,
                             'start_step1':start_step1,'step1_block_reason':step1_block_reason})
        return {'source_saved':True,'step1_task':{'id':'test-step1-task'} if start_step1 else None}
    async def start(self):pass
    async def stop(self):pass
    def status(self,p):return {'tasks':[],'runs':[],'queue':[],'artifacts':[],'latest_delivery':None}


@pytest.fixture
def api(tmp_path):
    schema='test_api_'+uuid.uuid4().hex
    conn=psycopg.connect('postgresql:///branch_agent_local',autocommit=True)
    conn.execute(f'CREATE SCHEMA "{schema}"')
    store=Store(f'postgresql:///branch_agent_local?options=-csearch_path%3D{schema}',tmp_path/'blobs')
    app=create_app(store=store,engine=IdleEngine(),data_dir=tmp_path)
    with TestClient(app) as client:
        password=(tmp_path/'local-login.txt').read_text().splitlines()[1].split('：',1)[1]
        login=client.post(BASE+'/auth/login',json={'username':'developer','password':password})
        assert login.status_code==200
        identity=login.json();client.headers['X-CSRF-Token']=identity['csrf_token']
        client.headers['Idempotency-Key']=str(uuid.uuid4())
        yield client,store,identity
    store.close();conn.execute(f'DROP SCHEMA "{schema}" CASCADE');conn.close()


def test_project_creation_idempotency_and_csrf(api):
    client,store,identity=api
    first=client.post(BASE+'/projects',json={'title':'验收测试'});assert first.status_code==200,first.text
    second=client.post(BASE+'/projects',json={'title':'验收测试'});assert second.json()['id']==first.json()['id']
    assert client.post(BASE+'/projects',json={'title':'不同正文'}).status_code==409
    client.headers.pop('X-CSRF-Token')
    assert client.post(BASE+'/projects',json={'title':'blocked'}).status_code==403


def test_artifact_defaults_seed_new_projects_and_project_edits_share_version_chain(api):
    client,store,_=api
    client.headers['Idempotency-Key']=str(uuid.uuid4())
    older=client.post(BASE+'/projects',json={'title':'旧项目'}).json()['id']
    defaults=client.get(BASE+'/artifact-defaults').json()
    assert defaults['adaptation_strategy']['result_kind']=='ready'
    assert defaults['adaptation_plan'] is None
    plan=deepcopy(defaults['starters']['adaptation_plan'])
    plan['payload']['title']='新项目的默认方案'
    saved=client.put(BASE+'/artifact-defaults',json={'adaptation_plan':plan})
    assert saved.status_code==200,saved.text
    client.headers['Idempotency-Key']=str(uuid.uuid4())
    newer=client.post(BASE+'/projects',json={'title':'新项目'}).json()['id']

    old_items=client.get(BASE+f'/projects/{older}/artifacts?kinds=adaptation_plan').json()['artifacts']
    old_plan=next(item for item in old_items if item['kind']=='adaptation_plan')
    assert old_plan['version'] is None
    old_write=client.put(BASE+f'/projects/{older}/artifacts/adaptation_plan',json={
        'content':plan,'expected_version':None})
    assert old_write.status_code==200,old_write.text
    assert old_write.json()['version']==old_write.json()['effective_version']==1
    items=client.get(BASE+f'/projects/{newer}/artifacts').json()['artifacts']
    seeded=next(item for item in items if item['kind']=='adaptation_plan')
    assert seeded['version']==seeded['effective_version']==1
    assert seeded['effective'] is True and seeded['origin']=='program'
    assert seeded['content'] is None and seeded['has_content'] is True
    assert seeded['stage']==4 and seeded['agent_key'] is None
    assert seeded['agent_name']=='Harness 写入'
    selected=client.get(BASE+f'/projects/{newer}/artifacts?kinds=adaptation_strategy,adaptation_plan').json()['artifacts']
    selected_plan=next(item for item in selected if item['kind']=='adaptation_plan')
    assert selected_plan['content']['payload']['title']=='新项目的默认方案'

    edited=deepcopy(selected_plan['content'])
    edited['payload']['title']='项目自己的修改'
    write=client.put(BASE+f'/projects/{newer}/artifacts/adaptation_plan',json={
        'content':edited,'expected_version':1})
    assert write.status_code==200,write.text
    assert write.json()['version']==write.json()['effective_version']==2
    assert write.json()['origin']=='user'
    assert write.json()['agent_name']=='人工编辑'
    assert write.json()['stage']==4
    assert client.put(BASE+f'/projects/{newer}/artifacts/adaptation_plan',json={
        'content':edited,'expected_version':1}).status_code==409
    assert client.get(BASE+'/artifact-defaults').json()['adaptation_plan']['payload']['title']=='新项目的默认方案'


def test_account_artifact_default_rejects_project_fixed_references(api):
    client,_,_=api
    plan=client.get(BASE+'/artifact-defaults').json()['starters']['adaptation_plan']
    plan['payload']['strategy_ref']={'record_id':str(uuid.uuid4()),'version':'1',
                                     'item_id':None,'json_pointer':None}
    rejected=client.put(BASE+'/artifact-defaults',json={'adaptation_plan':plan})
    assert rejected.status_code==400


def test_program_artifact_metadata_uses_actual_stage_and_harness_origin(api):
    from branch_agent.workflow import Workflow

    client,store,_=api
    project_id=client.post(BASE+'/projects',json={'title':'来源阶段测试'}).json()['id']
    plan=client.get(BASE+'/artifact-defaults').json()['starters']['adaptation_plan']
    assert client.put(BASE+f'/projects/{project_id}/artifacts/adaptation_plan',json={
        'content':plan,'expected_version':None}).status_code==200
    config=client.app.state.config
    workflow=Workflow(store)
    with store.transaction():
        graph=workflow.save(project_id,'nexo_graph',{
            'id':project_id,'name':'测试','description':'','prompt':'','revision':0,
            'updatedAt':'2026-09-26T00:00:00Z','chapters':[],
            'chapterEdges':[],'variables':[],'scenes':[],
        },stage=11,origin='program',config=config.resolve(project_id,'step11'),effective=True)
        events=workflow.save(project_id,'game_event_view',{
            'result_kind':'ready','payload':{'events':[],'event_links':[],'source_coverage':[]},
            'questions':[],'evidence_refs':[],'notes':[],
        },stage=6,origin='model',config=config.resolve(project_id,'step6'),effective=True)
        workflow.writeback_plan(project_id,events,6)

    workspace=client.app.state.artifact_workspace
    final=workspace.item(project_id,'nexo_graph')
    assert final['version']==graph['version']
    assert final['stage']==11 and final['agent_name']=='Harness 写入'
    assert final['agent_key'] is None
    written_plan=workspace.item(project_id,'adaptation_plan')
    assert written_plan['version']==2
    assert written_plan['stage']==6 and written_plan['agent_name']=='Harness 写入'


def test_account_and_blob_isolation(api):
    client,store,identity=api
    own=client.post(BASE+'/projects',json={'title':'我的项目'}).json();pid=own['id']
    other=store.put(new_record('project',None,title='私有',owner_account_id='another-account'))
    assert client.get(BASE+'/projects/'+other['id']).status_code==404
    blob=store.blob_put(other['id'],b'private','text/plain','private.txt')
    assert client.get(BASE+f'/projects/{pid}/blobs/{blob["id"]}/content').status_code==404
    assert client.get(BASE+f'/projects/{pid}/records/{other["id"]}').status_code==404
    assert all(p['id']!=other['id'] for p in client.get(BASE+'/projects').json()['items'])


def test_run_debug_shows_actual_input_and_model_output(api):
    client,store,_=api
    pid=client.post(BASE+'/projects',json={'title':'Run 调试'}).json()['id']
    client.headers['Idempotency-Key']=str(uuid.uuid4())
    cid=client.post(BASE+f'/projects/{pid}/conversations',json={'title':'会话'}).json()['id']
    config=client.app.state.config.resolve(pid,'step1')
    session=store.put(new_record('work_session',pid,conversation_id=cid,session_key='step1'))
    prompt=store.put(new_record('history_record',pid,conversation_id=cid,sequence=1,
        role='user',content={'storage':'inline_text','text':'开始改编'}))
    task=store.put(new_record('task',pid,conversation_id=cid,requested_by_message_id=prompt['id'],intent='generate',scope={
        'stage':1,'chapter_ids':[],'branch_ids':[],'target_refs':[],'description':'切片'}))
    run=store.put(new_record('run',pid,task_id=task['id'],agent_key='source_parser',
                             session_id=session['id'],config_version_id=config['id']))
    values={'instructions':{'storage':'inline_text','text':'切分原作'},
            'input_items':{'storage':'inline_json','value':[{'role':'user','content':'原作正文'}]},
            'tool_definitions':{'storage':'inline_json','value':[]}}
    from branch_agent.records import canonical_bytes
    import hashlib
    snapshot=store.put(new_record('context_snapshot',pid,task_id=task['id'],run_id=run['id'],
        session_id=session['id'],config_version_id=config['id'],model='gpt-5.6-luna',
        reasoning_effort='low',input_token_budget=10000,input_token_estimate=10,
        content_sha256=hashlib.sha256(canonical_bytes(values)).hexdigest(),**values))
    history=store.put(new_record('history_record',pid,conversation_id=cid,sequence=2,
        role='assistant',visibility='internal',kind='model_output',run_id=run['id'],
        content={'storage':'inline_json','value':{'output':[{'type':'message','content':[{'type':'output_text','text':'切片完成'}]}]}}))
    call=store.put(new_record('model_call',pid,task_id=task['id'],run_id=run['id'],
        operation_id=str(uuid.uuid4()),context_snapshot_id=snapshot['id'],
        response_history_ids=[history['id']],state='succeeded'))
    listing=client.get(BASE+f'/projects/{pid}/run-debug')
    assert listing.status_code==200,listing.text
    assert listing.json()['items'][0]['run']['id']==run['id']
    assert listing.json()['items'][0]['agent_name']==config['values']['prompts']['agent_names']['source_parser']
    assert listing.json()['items'][0]['model']=='gpt-5.6-luna'
    detail=client.get(BASE+f'/projects/{pid}/run-debug/{run["id"]}')
    assert detail.status_code==200,detail.text
    assert detail.json()['agent_name']==config['values']['prompts']['agent_names']['source_parser']
    assert detail.json()['snapshots'][0]['input_items']['value'][0]['content']=='原作正文'
    assert detail.json()['outputs'][0]['history']['content']['value']['output'][0]['content'][0]['text']=='切片完成'
    assert detail.json()['calls'][0]['id']==call['id']
    assert detail.json()['trace_spans']==[]
    store._connection().execute('''INSERT INTO sdk_trace_spans
        (project_id,run_id,trace_id,span_id,kind,name,started_at,metadata)
        VALUES (%s,%s,%s,%s,%s,%s,clock_timestamp(),%s)''',
        (pid,run['id'],'trace_test','span_test','function','get_artifact',Jsonb({})))
    traced=client.get(BASE+f'/projects/{pid}/run-debug/{run["id"]}')
    assert traced.status_code==200
    assert traced.json()['trace_spans'][0]['name']=='get_artifact'


def test_config_publication_and_old_snapshot_are_independent(api):
    client,store,identity=api
    pid=client.post(BASE+'/projects',json={'title':'配置测试'}).json()['id']
    service=client.app.state.config;old=service.resolve(pid,'step1')
    client.headers['Idempotency-Key']=str(uuid.uuid4())
    draft=client.post(BASE+f'/projects/{pid}/config',json={'values':{'model':{'reasoning_effort':'low'}},'scope_kind':'project'});assert draft.status_code==200,draft.text
    client.headers['Idempotency-Key']=str(uuid.uuid4())
    published=client.post(BASE+f'/projects/{pid}/config/{draft.json()["id"]}/publish',json={});assert published.status_code==200,published.text
    new=service.resolve(pid,'step1')
    assert new['values']['model']['reasoning_effort']=='low'
    assert old['values']['model']['reasoning_effort']=='high'
    assert new['values']['context']['input_token_cap']==1050000
    assert old['id']!=new['id']
    assert store.get(old['id'],pid)['values']==old['values']


def test_account_config_endpoints_apply_to_all_owned_projects(api):
    client,store,identity=api
    first=client.post(BASE+'/projects',json={'title':'账号配置一'}).json()['id']
    client.headers['Idempotency-Key']=str(uuid.uuid4())
    second=client.post(BASE+'/projects',json={'title':'账号配置二'}).json()['id']
    client.headers['Idempotency-Key']=str(uuid.uuid4())

    view=client.get(BASE+'/account/config?stage=step1')
    assert view.status_code==200,view.text
    draft=client.post(BASE+'/account/config',json={
        'values':{'model':{'reasoning_effort':'low'},
                  'context':{'step1_source':{'trigger_tokens':250000,'window_tokens':120000}}},
        'scope_kind':'project'
    })
    assert draft.status_code==200,draft.text
    client.headers['Idempotency-Key']=str(uuid.uuid4())
    published=client.post(BASE+f'/account/config/{draft.json()["id"]}/publish',json={})
    assert published.status_code==200,published.text

    service=client.app.state.config
    assert service.resolve(first,'step1')['values']['model']['reasoning_effort']=='low'
    assert service.resolve(second,'step1')['values']['model']['reasoning_effort']=='low'
    assert service.resolve(first,'step1')['values']['context']['step1_source']['trigger_tokens']==250000
    assert service.resolve(second,'step1')['values']['context']['step1_source']['window_tokens']==120000

    client.headers['Idempotency-Key']=str(uuid.uuid4())
    created=client.post(BASE+f'/projects/{first}/conversations',json={'title':'导入配置验证'})
    assert created.status_code==200,created.text
    client.headers['Idempotency-Key']=str(uuid.uuid4())
    imported=client.post(BASE+f'/projects/{first}/source',json={
        'conversation_id':created.json()['id'],'text':'用于验证全局原作上限的简短原文。'
    })
    assert imported.status_code==200,imported.text
    assert imported.json()['source_window_threshold']==250000
    assert imported.json()['source_window_tokens']==120000
    assert imported.json()['source_mode']=='full_text'
    assert imported.json()['step1_started'] is True
    assert imported.json()['step1_task_id']=='test-step1-task'
    assert client.app.state.engine.imports[-1]['start_step1'] is True


def test_source_import_reports_independent_step1_branch_modes(api):
    client,_,_=api
    project=client.post(BASE+'/projects',json={'title':'双分支预估'}).json()['id']
    client.headers['Idempotency-Key']=str(uuid.uuid4())
    conversation=client.post(BASE+f'/projects/{project}/conversations',json={'title':'原作导入'}).json()['id']
    client.headers['Idempotency-Key']=str(uuid.uuid4())
    draft=client.post(BASE+'/account/config',json={
        'values':{'context':{'step1_source':{'trigger_tokens':1}}},
        'scope_kind':'agent','scope_key':'source_character_parser',
    })
    assert draft.status_code==200,draft.text
    client.headers['Idempotency-Key']=str(uuid.uuid4())
    published=client.post(BASE+f'/account/config/{draft.json()["id"]}/publish',json={})
    assert published.status_code==200,published.text
    client.headers['Idempotency-Key']=str(uuid.uuid4())
    imported=client.post(BASE+f'/projects/{project}/source',json={
        'conversation_id':conversation,'text':'韩立走出山村，踏上修行之路。',
    })
    assert imported.status_code==200,imported.text
    receipt=imported.json()
    assert receipt['source_mode']=='mixed'
    assert receipt['step1_branches']['global']['source_mode']=='full_text'
    assert receipt['step1_branches']['character']['source_mode']=='sliding_window'
    assert receipt['step1_branches']['global']['agent_key']=='source_global_parser'
    assert receipt['step1_branches']['character']['agent_key']=='source_character_parser'
    assert receipt['step1_branches']['character']['source_window_threshold']==1
    assert receipt['step1_branches']['global']['input_budget']>receipt['step1_branches']['global']['source_tokens']
    assert receipt['step1_started'] is True


def test_source_import_saves_but_does_not_start_when_step1_budget_is_insufficient(api,monkeypatch):
    import branch_agent.app as app_module
    client,_,_=api
    project,conversation=create_space(client)
    monkeypatch.setattr(app_module,'input_budget',lambda _values:1)
    result=client.post(BASE+f'/projects/{project}/source',json={
        'conversation_id':conversation,'text':'预算不足时仍保留的原作。'})
    assert result.status_code==200,result.text
    receipt=result.json()
    assert receipt['admitted'] is False
    assert receipt['step1_started'] is False
    assert receipt['step1_task_id'] is None
    assert receipt['step1_reason']=='budget_insufficient'
    assert client.app.state.engine.imports[-1]['start_step1'] is False
    assert client.app.state.engine.imports[-1]['step1_block_reason']=='budget_insufficient'


def test_conversation_activity_is_scoped_ordered_and_tail_bounded(api):
    client,store,_=api
    project,conversation=create_space(client)
    client.headers['Idempotency-Key']=str(uuid.uuid4())
    other=client.post(BASE+f'/projects/{project}/conversations',json={'title':'另一个会话'}).json()['id']
    names=['chat.activity','chat.reply.started','chat.reply.delta','chat.reply.completed',
           'chat.reply.failed','tool.called','chat.activity']
    targets=[conversation,conversation,other,conversation,conversation,conversation,conversation]
    for sequence,(name,target) in enumerate(zip(names,targets),1):
        store.put(new_record('runtime_event',project,sequence=sequence,event_name=name,
                             conversation_id=target,payload={'text':str(sequence)}))
    endpoint=BASE+f'/projects/{project}/conversations/{conversation}/activity'
    tail=client.get(endpoint+'?limit=2').json()
    assert [event['sequence'] for event in tail['events']]==[5,7]
    assert tail['last_sequence']==7
    assert tail['next_sequence']==7
    assert tail['truncated_before'] is True
    older=client.get(endpoint+'?before=5&limit=2').json()
    assert [event['sequence'] for event in older['events']]==[2,4]
    assert older['truncated_before'] is True
    first=client.get(endpoint+'?tail=false&after=0&limit=2').json()
    assert [event['sequence'] for event in first['events']]==[1,2]
    assert first['has_more'] is True
    second=client.get(endpoint+'?after=2&limit=2').json()
    assert [event['sequence'] for event in second['events']]==[4,5]
    assert client.get(BASE+f'/projects/{project}/conversations/{uuid.uuid4()}/activity').status_code==404


def test_account_config_schema_validation_is_read_only_and_reports_strict_errors(api):
    client,store,identity=api
    view=client.get(BASE+'/account/config?stage=step5').json()
    schemas=deepcopy(view['schemas'])
    schemas['game_event_view']['description']='调试同学上传后的自定义说明'

    client.headers['Idempotency-Key']=str(uuid.uuid4())
    checked=client.post(BASE+'/account/config/validate',json={
        'values':view['values'],'schemas':schemas,'stage':'step5',
    })
    assert checked.status_code==200,checked.text
    assert checked.json()['valid'] is True
    assert checked.json()['enabled'] is True
    assert checked.json()['schema_id']=='game_event_view'
    assert checked.json()['schema_sha256']
    assert not store.list(None,'config_version',limit=10)

    invalid=deepcopy(schemas)
    invalid['game_event_view'].pop('additionalProperties')
    client.headers['Idempotency-Key']=str(uuid.uuid4())
    rejected=client.post(BASE+'/account/config/validate',json={
        'values':view['values'],'schemas':invalid,'stage':'step5',
    })
    assert rejected.status_code==400
    assert 'additionalProperties=false' in rejected.json()['error']['message']


def test_unsaved_instruction_layers_have_server_preview_without_publication(api):
    client, store, _ = api
    values = client.get(BASE+'/account/config?stage=step5').json()['values']
    values['prompts']['base'] = '未发布的创作指令'
    values['prompts']['harness']['stages']['step5'] = '未发布的阶段协议'
    result = client.post(BASE+'/account/config/instructions-preview', json={
        'stage':'step5', 'values':values})
    assert result.status_code == 200, result.text
    shown = result.json()
    assert shown['parts']['creative_base'] == '未发布的创作指令'
    assert shown['parts']['harness_stage'] == '未发布的阶段协议'
    assert shown['final'].endswith(values['prompts']['harness']['ask_user'])
    assert not store.list(None,'config_version',limit=10)

    source_values=client.get(BASE+'/account/config?stage=step1').json()['values']
    source_preview=client.post(BASE+'/account/config/instructions-preview',json={
        'stage':'step1','values':source_values})
    assert source_preview.status_code==200,source_preview.text
    branches=source_preview.json()['views']
    assert '只负责作品事件视图及其分析' in branches['global']['final']
    assert '人物分支不输出顶层 analysis' in branches['character']['final']
    assert '滑动窗口固定协议' not in branches['global']['final']
    assert branches['global']['parts']['tool_guidance']==source_values['prompts']['harness']['no_ask_user']
    window_preview=client.post(BASE+'/account/config/instructions-preview',json={
        'stage':'step1','values':source_values,'step1_mode':'window','step1_view':'character'})
    assert window_preview.status_code==200,window_preview.text
    assert '滑动窗口固定协议' in window_preview.json()['final']
    assert '当前独立产物：主要人物事件' in window_preview.json()['final']
    assert '也不输出 analysis' in window_preview.json()['final']


def test_summary_agent_configuration_preview_and_publication(api):
    client, _, _ = api
    project = client.post(BASE+'/projects', json={'title':'摘要配置验证'}).json()['id']
    current = client.get(BASE+'/account/config?stage=aux.summary')
    assert current.status_code == 200, current.text
    assert current.json()['selected_agent'] == 'context_summarizer'

    values = deepcopy(current.json()['values'])
    values['prompts']['summary'] = '只归纳已发生的对话与待办。'
    values['summary']['target_tokens'] = 900
    preview = client.post(BASE+'/account/config/instructions-preview', json={
        'stage':'aux.summary', 'values':values})
    assert preview.status_code == 200, preview.text
    assert preview.json()['parts']['creative_stage'] == values['prompts']['summary']
    assert '只归纳已发生的对话与待办' in preview.json()['final']
    checked = client.post(BASE+'/account/config/validate', json={
        'stage':'aux.summary', 'values':values, 'schemas':current.json()['schemas']})
    assert checked.status_code == 200, checked.text
    assert checked.json()['enabled'] is False

    client.headers['Idempotency-Key'] = str(uuid.uuid4())
    draft = client.post(BASE+'/account/config', json={
        'scope_kind':'auxiliary', 'scope_key':'aux.summary',
        'values':{'prompts':{'layout_version':3,'summary':values['prompts']['summary']},
                  'summary':{'target_tokens':900}}})
    assert draft.status_code == 200, draft.text
    client.headers['Idempotency-Key'] = str(uuid.uuid4())
    published = client.post(BASE+f'/account/config/{draft.json()["id"]}/publish', json={})
    assert published.status_code == 200, published.text
    summary = client.app.state.config.resolve(project, 'coordinator')['values']['auxiliary_configs']['aux.summary']
    assert summary['summary']['target_tokens'] == 900
    assert summary['prompts']['summary'] == values['prompts']['summary']
    assert summary['prompts']['stage_agents']['aux.summary'] == 'context_summarizer'


def test_format_repair_agent_configuration_preview_and_publication(api):
    client, _, _ = api
    project = client.post(BASE+'/projects', json={'title':'格式修复配置验证'}).json()['id']
    current = client.get(BASE+'/account/config?stage=aux.format_repair')
    assert current.status_code == 200, current.text
    assert current.json()['selected_agent'] == 'format_repairer'
    values = deepcopy(current.json()['values'])
    values['prompts']['stage'] = '按精确错误做指定位置的格式修正。'
    preview = client.post(BASE+'/account/config/instructions-preview', json={
        'stage':'aux.format_repair', 'values':values})
    assert preview.status_code == 200, preview.text
    assert preview.json()['parts']['creative_stage'] == values['prompts']['stage']
    assert 'candidate_sha256' in preview.json()['final']
    assert 'tool_guidance' not in preview.json()['parts']
    checked = client.post(BASE+'/account/config/validate', json={
        'stage':'aux.format_repair', 'values':values, 'schemas':current.json()['schemas']})
    assert checked.status_code == 200, checked.text
    assert checked.json()['enabled'] is False

    client.headers['Idempotency-Key'] = str(uuid.uuid4())
    draft = client.post(BASE+'/account/config', json={
        'scope_kind':'auxiliary', 'scope_key':'aux.format_repair',
        'values':{'prompts':{'layout_version':3,'stage':values['prompts']['stage']},
                  'model':{'reasoning_effort':'low'}}})
    assert draft.status_code == 200, draft.text
    client.headers['Idempotency-Key'] = str(uuid.uuid4())
    published = client.post(BASE+f'/account/config/{draft.json()["id"]}/publish', json={})
    assert published.status_code == 200, published.text
    repair = client.app.state.config.resolve(project, 'coordinator')['values']['auxiliary_configs']['aux.format_repair']
    assert repair['prompts']['stage'] == values['prompts']['stage']
    assert repair['model']['reasoning_effort'] == 'low'


def test_account_editor_autosave_keeps_incomplete_schema_without_publishing(api):
    client,store,identity=api
    original=client.get(BASE+'/account/config?stage=step5').json()
    before=client.get(BASE+'/account/config/editor?scope_kind=stage&scope_key=step5')
    assert before.status_code==200 and before.json()['revision']==0
    working={'values':original['values'],'base_values':original['values'],
             'schemas':original['schemas'],'schema_texts':{'game_event_view':'{"type":'},
             'raw_fields':{'model.max_output_tokens':''}}
    saved=client.put(BASE+'/account/config/editor',json={'scope_kind':'stage','scope_key':'step5',
      'expected_revision':0,'payload':working})
    assert saved.status_code==200,saved.text
    revision=saved.json()['revision']
    restored=client.get(BASE+'/account/config/editor?scope_kind=stage&scope_key=step5')
    assert restored.status_code==200
    assert restored.json()['payload']['schema_texts']['game_event_view']=='{"type":'
    assert restored.json()['payload']['raw_fields']['model.max_output_tokens']==''
    assert client.get(BASE+'/account/config?stage=step5').json()['schemas']==original['schemas']
    assert not store.list(None,'config_version',limit=10)
    conflict=client.put(BASE+'/account/config/editor',json={'scope_kind':'stage','scope_key':'step5',
      'expected_revision':0,'payload':working})
    assert conflict.status_code==409
    cleared=client.request('DELETE',BASE+'/account/config/editor',json={
      'scope_kind':'stage','scope_key':'step5','expected_revision':revision})
    assert cleared.status_code==200
    assert client.get(BASE+'/account/config/editor?scope_kind=stage&scope_key=step5').json()['payload'] is None


def test_invalid_publication_keeps_previous_published_config(api):
    client,store,_=api
    original=client.get(BASE+'/account/config?stage=step1').json()
    client.headers['Idempotency-Key']=str(uuid.uuid4())
    draft=client.post(BASE+'/account/config',json={
      'scope_kind':'project','values':{'model':{'max_output_tokens':999999999}}})
    assert draft.status_code==200,draft.text
    client.headers['Idempotency-Key']=str(uuid.uuid4())
    refused=client.post(BASE+f'/account/config/{draft.json()["id"]}/publish',json={})
    assert refused.status_code==400
    current=client.get(BASE+'/account/config?stage=step1').json()
    assert current['values']['model']['max_output_tokens']==original['values']['model']['max_output_tokens']
    assert not [version for version in current['versions'] if version['state']=='published']


def test_permanent_project_delete_requires_confirmation_and_is_idempotent(api):
    client,store,identity=api
    pid,cid=create_space(client)
    blob=store.blob_put(pid,b'private source','text/plain','source.txt')
    blob_path=store.blob_dir/blob['storage_key']
    endpoint=BASE+f'/projects/{pid}'

    client.headers['Idempotency-Key']=str(uuid.uuid4())
    refused=client.request('DELETE',endpoint,json={'confirm':False})
    assert refused.status_code==400

    key=str(uuid.uuid4());client.headers['Idempotency-Key']=key
    deleted=client.request('DELETE',endpoint,json={'confirm':True})
    assert deleted.status_code==200,deleted.text
    assert deleted.json()=={'deleted':True,'project_id':pid}
    assert not blob_path.exists()
    assert store.get(pid,pid) is None
    assert store.get(cid,pid) is None
    assert all(item['id']!=pid for item in client.get(BASE+'/projects').json()['items'])

    repeated=client.request('DELETE',endpoint,json={'confirm':True})
    assert repeated.status_code==200
    assert repeated.json()==deleted.json()


def create_space(client):
    client.headers['Idempotency-Key']=str(uuid.uuid4())
    response=client.post(BASE+'/projects',json={'title':'API边界测试'})
    assert response.status_code==200,response.text
    pid=response.json()['id']
    client.headers['Idempotency-Key']=str(uuid.uuid4())
    response=client.post(BASE+f'/projects/{pid}/conversations',json={'title':'测试会话'})
    assert response.status_code==200,response.text
    return pid,response.json()['id']


def test_bad_inputs_origin_and_cookie_are_structured_errors(api):
    client,store,identity=api
    assert client.post(BASE+'/auth/login',json={'username':[],'password':{}}).status_code==400
    assert client.post(BASE+'/projects',json=[]).status_code==400
    assert client.post(BASE+'/projects',json={'title':23}).status_code==400
    assert client.post(BASE+'/projects',json={'title':'cross origin'},headers={'Origin':'https://untrusted.example'}).status_code==403
    client.headers['Idempotency-Key']='x'*129
    assert client.post(BASE+'/projects',json={'title':'long key'}).status_code==400
    assert client.get(BASE+'/projects/not-a-uuid').status_code==404
    assert client.get(BASE+'/projects?cursor=not-a-number').json()['error']['code']=='invalid_request'
    client.cookies.clear()
    client.cookies.set('branch_session',client.app.state.auth.signer.dumps(['malformed signed payload']))
    assert client.get(BASE+'/auth/me').status_code==401


def test_workflow_blocked_is_readable_conflict(api):
    from branch_agent.workflow import WorkflowBlocked
    client,store,identity=api;pid,cid=create_space(client)
    def blocked(*args,**kwargs):raise WorkflowBlocked('confirmation_required',{'stage':4})
    client.app.state.engine.submit_message=blocked
    response=client.post(BASE+f'/projects/{pid}/conversations/{cid}/messages',json={'text':'继续','mode':'queue'})
    assert response.status_code==409,response.text
    assert response.json()['error']['code']=='confirmation_required'
    assert '确认' in response.json()['error']['message']
    assert response.json()['error']['details']=={'stage':4}


def test_docx_body_order_nested_tables_and_idempotent_import(api):
    from docx import Document
    client,store,identity=api;pid,cid=create_space(client)
    doc=Document();doc.add_paragraph('开场段落')
    table=doc.add_table(rows=1,cols=1);cell=table.cell(0,0);cell.text='表格第一句'
    inner=cell.add_table(rows=1,cols=1);inner.cell(0,0).text='嵌套表格'
    cell.add_paragraph('表格后一句');doc.add_paragraph('结尾段落')
    stream=BytesIO();doc.save(stream);raw=stream.getvalue()
    key=str(uuid.uuid4());client.headers['Idempotency-Key']=key
    endpoint=BASE+f'/projects/{pid}/source'
    response=client.post(endpoint,data={'conversation_id':cid},files={'file':('原作.docx',raw,'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert response.status_code==200,response.text
    text=client.app.state.engine.imports[-1]['text']
    phrases=['开场段落','表格第一句','嵌套表格','表格后一句','结尾段落']
    assert [text.index(x) for x in phrases]==sorted(text.index(x) for x in phrases)
    repeated=client.post(endpoint,data={'conversation_id':cid},files={'file':('原作.docx',raw)})
    assert repeated.status_code==200 and len(client.app.state.engine.imports)==1
    assert repeated.json()['attachment']['id']==response.json()['attachment']['id']
    assert client.post(endpoint,data={'conversation_id':cid},files={'file':('原作.docx',b'broken archive')}).status_code==400
    assert client.post(endpoint,data={'conversation_id':cid}).status_code==400


def test_json_and_file_source_share_byte_limit_and_invalid_inputs(api,monkeypatch):
    import branch_agent.app as app_module
    client,store,identity=api;pid,cid=create_space(client)
    endpoint=BASE+f'/projects/{pid}/source';monkeypatch.setattr(app_module,'MAX_SOURCE_BYTES',10)
    assert client.post(endpoint,json={'conversation_id':cid,'text':'中文中文'}).status_code==413
    assert client.post(endpoint,data={'conversation_id':cid},files={'file':('原作.txt','中文中文'.encode())}).status_code==413
    assert client.post(endpoint,json={'conversation_id':cid,'text':23}).status_code==400
    assert client.post(endpoint,data={'conversation_id':cid},files={'file':('原作.txt',b'\xff')}).status_code==400
    assert not client.app.state.engine.imports


def test_utf16_le_source_import(api):
    client,store,identity=api;pid,cid=create_space(client)
    original='第一章\n韩立走进山村。'
    raw=original.encode('utf-16')
    response=client.post(BASE+f'/projects/{pid}/source',data={'conversation_id':cid},
        files={'file':('原作.txt',raw,'text/plain')})
    assert response.status_code==200,response.text
    assert client.app.state.engine.imports[-1]['text']==original
    attachment=response.json()['attachment']
    assert store.blob_read(attachment['id'],pid)==raw


def test_release_endpoint_reports_update_without_caching(api):
    client,store,identity=api
    response=client.get(BASE+'/release')
    assert response.status_code==200
    assert response.headers['cache-control']=='no-store'
    assert response.json()['id']
    assert response.json()['title']
    assert response.json()['changes']


def test_blob_json_public_projection_and_binary_download(api):
    client,store,identity=api;pid,cid=create_space(client)
    data={'instructions':'ordinary content','encrypted_content':'not-visible','nested':{'api_key':'secret','Authorization':'Bearer token-private'},'reasoning_tokens':23}
    blob=store.blob_put(pid,json.dumps(data).encode(),'application/json','context.json')
    response=client.get(BASE+f'/projects/{pid}/blobs/{blob["id"]}/content')
    assert response.status_code==200
    assert 'secret' not in response.text and 'not-visible' not in response.text and 'token-private' not in response.text
    assert '已遮蔽' in response.json()['nested']['Authorization']
    assert response.json()['reasoning_tokens']==23
    assert response.headers['X-Content-Projection']=='public-redacted'
    html=store.blob_put(pid,b'<script>alert(1)</script>','text/html','active.html')
    assert client.get(BASE+f'/projects/{pid}/blobs/{html["id"]}/content').headers['content-type'].startswith('text/plain')
    binary=store.blob_put(pid,b'PKbinary','application/octet-stream','a.docx')
    assert 'attachment' in client.get(BASE+f'/projects/{pid}/blobs/{binary["id"]}/content').headers['content-disposition']


def test_relations_use_fixed_projection_not_uuid_text_substrings(api):
    from branch_agent.workflow import Workflow,ref
    client,store,identity=api;pid,cid=create_space(client)
    workflow=Workflow(store)
    first=workflow.save(pid,'source_text','第一版',origin='import',effective=True)
    second=workflow.save(pid,'source_text','第二版',origin='import',effective=True)
    event=store.put(new_record('runtime_event',pid,sequence=1,event_name='fixed.reference',subject_ref=ref(first),payload={}))
    decoy=store.put(new_record('history_record',pid,conversation_id=cid,sequence=1,role='user',content={'storage':'inline_text','text':'只是正文提到了 '+first['id']}))
    response=client.get(BASE+f'/projects/{pid}/records/{event["id"]}/relations?direction=out')
    assert response.status_code==200,response.text
    rows=response.json()['items']
    link=next(x for x in rows if x['target']==first['id'])
    assert link['target_ref']['version']=='1'
    assert all(x['target']!=second['id'] for x in rows)
    incoming=client.get(BASE+f'/projects/{pid}/records/{first["id"]}/relations?direction=in').json()['items']
    assert any(x['source']==event['id'] for x in incoming)
    assert all(x['source']!=decoy['id'] for x in incoming)


def test_history_projects_knowledge_asset_to_readable_model_content(api):
    from branch_agent.workflow import Workflow, ref
    client,store,identity=api;pid,cid=create_space(client)
    workflow=Workflow(store)
    source=workflow.save(pid,'source_text','甲在雾港寻找失踪的妹妹。',origin='import',effective=True)
    from test_runtime import step2_response
    result=step2_response('甲必须在风暴抵达前找到妹妹。','保留兄妹关系。')
    for key in ('source_global_events_ref','source_global_analysis_ref','source_character_events_ref'):
        result['payload'][key]=ref(source)
    result['notes']=['保留兄妹关系。']
    version=workflow.save(pid,'source_knowledge_asset',result,stage=2)
    conversation=store.get(cid,pid);conversation['last_message_seq']=1;store.update(conversation,conversation['row_version'])
    message=store.put(new_record('history_record',pid,conversation_id=cid,sequence=1,role='assistant',
        visibility='conversation',content={'storage':'inline_text','text':'Step2 已生成草稿 v1，请查看后确认，或提出具体修改。'}))
    store.put(new_record('runtime_event',pid,sequence=1,event_name='artifact.presented',conversation_id=cid,
        payload={'artifact_ref':ref(version),'message_id':message['id'],'selections':[{'item_id':None,'json_pointer':''}]}))

    shown=client.get(BASE+f'/projects/{pid}/conversations/{cid}/history').json()['items'][0]

    assert shown['content']['text'].startswith('Step2 已生成草稿')
    assert 'Step 2 · 原作知识资产' in shown['display_text']
    assert '甲必须在风暴抵达前找到妹妹。' in shown['display_text']
    assert '保留兄妹关系。' in shown['display_text']
    assert '"payload"' not in shown['display_text']


def test_markdown_history_and_download_are_bound_to_the_presented_version(api):
    from branch_agent.workflow import Workflow, ref
    from test_runtime import step2_response
    client,store,_=api;pid,cid=create_space(client)
    workflow=Workflow(store)
    source=workflow.save(pid,'source_text','甲在雾港寻找妹妹。',origin='import',effective=True)
    result=step2_response('旧版前提。','保留兄妹关系。')
    for key in ('source_global_events_ref','source_global_analysis_ref','source_character_events_ref'):
        result['payload'][key]=ref(source)
    first=workflow.save(pid,'source_knowledge_asset',result,stage=2)
    conversation=store.get(cid,pid);conversation['last_message_seq']=1
    store.update(conversation,conversation['row_version'])
    message=store.put(new_record('history_record',pid,conversation_id=cid,sequence=1,role='assistant',
        visibility='conversation',content={'storage':'inline_text','text':'旧的展示文本'}))
    store.put(new_record('runtime_event',pid,sequence=1,event_name='artifact.presented',conversation_id=cid,
        payload={'artifact_ref':ref(first),'message_id':message['id'],'selections':[{'item_id':None,'json_pointer':''}]}))
    newer=deepcopy(result);newer['payload']['premise']='新版前提。'
    workflow.save(pid,'source_knowledge_asset',newer,stage=2)

    shown=client.get(BASE+f'/projects/{pid}/conversations/{cid}/history').json()['items'][0]
    assert shown['display_format']=='markdown'
    assert shown['artifact_ref']==ref(first)
    assert shown['markdown_download_url'].endswith(f'/artifacts/{first["artifact_id"]}/versions/1/download.md')
    download=client.get(shown['markdown_download_url'])
    assert download.status_code==200
    assert download.headers['content-type'].startswith('text/markdown')
    assert 'attachment;' in download.headers['content-disposition']
    assert download.headers['x-content-type-options']=='nosniff'
    assert download.text==shown['display_text']
    assert '# Step 2 · 原作知识资产（v1）' in download.text
    assert '旧版前提。' in download.text and '新版前提。' not in download.text
    assert '请确认以上结果' not in download.text
    assert client.get(BASE+f'/projects/{pid}/artifacts/{first["artifact_id"]}/versions/1/download').status_code==200


def test_event_source_excerpt_uses_markdown_display_without_artifact_download(api):
    client,store,_=api;pid,cid=create_space(client)
    conversation=store.get(cid,pid);conversation['last_message_seq']=1
    store.update(conversation,conversation['row_version'])
    excerpt='# 事件原文\n\n````text\n甲 <script>见乙</script>。\n````'
    message=store.put(new_record('history_record',pid,conversation_id=cid,sequence=1,role='assistant',
        visibility='conversation',content={'storage':'inline_text','text':excerpt}))
    store.put(new_record('runtime_event',pid,sequence=1,event_name='event_source.presented',
        conversation_id=cid,payload={'message_id':message['id']}))
    shown=client.get(BASE+f'/projects/{pid}/conversations/{cid}/history').json()['items'][0]
    assert shown['display_format']=='markdown'
    assert shown['content']['text']==excerpt
    assert 'markdown_download_url' not in shown


def test_history_records_and_sse_continue_after_ten_thousand(api):
    client,store,identity=api;pid,cid=create_space(client)
    connection=store._connection();count=10007
    with connection.transaction():
        with connection.cursor() as cursor:
            cursor.executemany('INSERT INTO history_records(data) VALUES(%s)',[(Jsonb(new_record('history_record',pid,conversation_id=cid,sequence=i,role='user',content={'storage':'inline_text','text':str(i)})),) for i in range(1,count+1)])
            cursor.executemany('INSERT INTO runtime_events(data) VALUES(%s)',[(Jsonb(new_record('runtime_event',pid,sequence=i,event_name='pagination.check',conversation_id=cid,payload={})),) for i in range(1,count+1)])
    history=client.get(BASE+f'/projects/{pid}/conversations/{cid}/history?tail=true&limit=3').json()
    assert [r['sequence'] for r in history['items']]==[count-2,count-1,count]
    assert history['start_cursor']==count-3
    after=client.get(BASE+f'/projects/{pid}/conversations/{cid}/history?after_sequence=10005').json()
    assert [r['sequence'] for r in after['items']]==[10006,10007]
    records=client.get(BASE+f'/projects/{pid}/records?record_type=history_record&cursor=10000&limit=100').json()
    assert len(records['items'])==7 and not records['has_more']
    events=client.app.state.event_rows(pid,10005)
    assert [r['sequence'] for r in events]==[10006,10007]
    assert client.get(BASE+f'/projects/{pid}/events',headers={'Last-Event-ID':'bad'}).status_code==400
    endpoint=next(route.endpoint for route in client.app.routes if getattr(route,'path','')==BASE+'/projects/{p}/events')
    class EventRequest:
        headers={'Last-Event-ID':'10006'}
        cookies={'branch_session':client.cookies.get('branch_session')}
        async def is_disconnected(self):return False
    async def first_event():
        response=await endpoint(EventRequest(),pid,after=0)
        iterator=response.body_iterator
        try:return await iterator.__anext__()
        finally:await iterator.aclose()
    assert 'id: 10007\n' in asyncio.run(first_event())
    assert client.get(BASE+f'/projects/{pid}/records?cursor=-1').status_code==400


def test_docx_unreadable_body_is_not_silently_imported(api):
    from docx import Document
    from docx.oxml import OxmlElement
    client,store,identity=api;pid,cid=create_space(client)
    doc=Document();doc.add_paragraph('可读段落')
    doc.add_paragraph().add_run()._r.append(OxmlElement('w:drawing'))
    stream=BytesIO();doc.save(stream)
    response=client.post(BASE+f'/projects/{pid}/source',data={'conversation_id':cid},files={'file':('原作.docx',stream.getvalue())})
    assert response.status_code==400 and '完整' in response.json()['error']['message']
    assert not client.app.state.engine.imports
    assert len(client.get(BASE+'/projects?limit=0').json()['items'])==1
    assert (client.app.state.auth.root/'local-auth.json').stat().st_mode&0o777==0o600
    assert (client.app.state.auth.root/'local-login.txt').stat().st_mode&0o777==0o600
